"""Regenerate the skill's assets/letterhead-template.docx. Presenter tool; lives outside the skill folder.

    python extras/build_letterhead.py

Canyon & Crest LLP is fictional. The template has {{placeholders}} that
scripts/fill_template.py replaces, and a {{BODY}} paragraph where the letter body goes.
"""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "engagement-letter" / "assets" / "letterhead-template.docx"
FONT = "Cambria"
GRAY = RGBColor(0x55, 0x55, 0x55)


def para(doc, text="", size=None, bold=False, align=None, color=None, after=0, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    return p


def bottom_rule(p):
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "8", "w:space": "6", "w:color": "333333"}.items():
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)


def page_field(run):
    for tag, extra in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None),
                       ("w:fldChar", {"w:fldCharType": "end"})):
        el = OxmlElement(tag)
        if extra:
            for k, v in extra.items():
                el.set(qn(k), v)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        run._r.append(el)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(1.1)
sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.9)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

body = doc.styles.add_style("Letter Body", WD_STYLE_TYPE.PARAGRAPH)
body.base_style = normal
body.paragraph_format.space_after = Pt(9)
body.paragraph_format.line_spacing = 1.1

heading = doc.styles.add_style("Letter Heading", WD_STYLE_TYPE.PARAGRAPH)
heading.base_style = body
heading.font.bold = True
heading.paragraph_format.space_before = Pt(6)
heading.paragraph_format.space_after = Pt(3)
heading.paragraph_format.keep_with_next = True

# Letterhead
para(doc, "CANYON & CREST LLP", size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, "ATTORNEYS AT LAW", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY, after=2)
rule = para(doc, "480 East Canyon Crest Road, Suite 300  ·  Provo, Utah 84604  ·  (801) 555-0147  ·  canyoncrest.example",
            size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY, after=24)
bottom_rule(rule)

# Date, confidentiality legend, client block, Re line
para(doc, "{{letter_date}}", after=18)
para(doc, "PRIVILEGED AND CONFIDENTIAL", size=9, bold=True)
para(doc, "ATTORNEY-CLIENT COMMUNICATION", size=9, after=12)
para(doc, "{{client_name}}")
para(doc, "{{client_address}}", after=18)
re_line = para(doc, "Re:\tEngagement of Canyon & Crest LLP: {{matter_description}}", bold=True, after=18)
re_line.paragraph_format.left_indent = Inches(0.5)
re_line.paragraph_format.first_line_indent = Inches(-0.5)

# Body goes here
para(doc, "{{BODY}}", style="Letter Body")

# Footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Canyon & Crest LLP  ·  Page ")
fr.font.size, fr.font.color.rgb = Pt(8), GRAY
pr = fp.add_run()
pr.font.size, pr.font.color.rgb = Pt(8), GRAY
page_field(pr)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Wrote {OUT}")
