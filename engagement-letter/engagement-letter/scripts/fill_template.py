#!/usr/bin/env python3
"""Fill the firm letterhead with the tailored engagement letter.

Usage:
    python fill_template.py --intake intake.json --clauses ../references/clauses-litigation.md \
        [--deadlines deadlines.json] --out "Engagement Letter.docx"

How the letter is assembled:
  1. Start with the firm's standard form (assets/standard-form.md).
  2. Each "## Section" in the clause file replaces the standard section with the
     same heading. A section with a new heading is inserted after the one above it.
  3. {{placeholders}} are filled from the intake JSON and, for litigation,
     the JSON produced by compute_deadline.py.
  4. The result is written into assets/letterhead-template.docx.

intake.json fields:
  client_name, matter_description, matter_type (litigation | transactional | flat-fee),
  fee_terms, retainer_amount; optional: client_address, letter_date (YYYY-MM-DD).

Requires python-docx (pip install python-docx).
"""

import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document

SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL_DIR / "assets" / "letterhead-template.docx"
STANDARD_FORM = SKILL_DIR / "assets" / "standard-form.md"

UNPRINTED_HEADINGS = {"Opening", "Closing"}
BODY_STYLE, HEADING_STYLE = "Letter Body", "Letter Heading"
MATTER_TYPES = ("litigation", "transactional", "flat-fee")
REQUIRED_FIELDS = ("client_name", "matter_description", "matter_type", "fee_terms", "retainer_amount")

SERVICE_METHOD_LABELS = {
    "personal": "personal delivery",
    "mail": "mail (the date you signed the receipt)",
    "email": "electronic acceptance of service",
    "out-of-state": "service outside Utah",
}


# --- Markdown form -> sections ------------------------------------------------

def parse_sections(md_text):
    """Return [(heading, [paragraph, ...]), ...]. A line ending in '\\' keeps its line break."""
    md_text = re.sub(r"<!--.*?-->", "", md_text, flags=re.S)
    sections, current = [], None
    for block in re.split(r"\n\s*\n", md_text):
        lines = [ln.rstrip() for ln in block.strip().splitlines()]
        if not lines or not lines[0]:
            continue
        if lines[0].startswith("## "):
            current = (lines[0][3:].strip(), [])
            sections.append(current)
            lines = lines[1:]
            if not lines:
                continue
        if current is None:
            continue  # text before the first heading (e.g. a title) is ignored
        text = ""
        for ln in lines:
            if ln.endswith("\\"):
                text += ln[:-1].rstrip() + "\n"
            else:
                text += ln.strip() + " "
        current[1].append(text.strip(" "))
    return sections


def merge_sections(standard, clauses):
    """Clause sections replace same-named standard sections; new ones go after the section above them."""
    merged = [(h, list(p)) for h, p in standard]
    anchor = None
    for heading, paras in clauses:
        names = [h for h, _ in merged]
        if heading in names:
            merged[names.index(heading)] = (heading, list(paras))
        else:
            pos = names.index(anchor) + 1 if anchor in names else len(merged) - 1
            merged.insert(pos, (heading, list(paras)))
        anchor = heading
    return merged


def fill(text, fields):
    missing = sorted({m for m in re.findall(r"{{(\w+)}}", text) if m not in fields})
    if missing:
        raise KeyError("Missing value(s) for placeholder(s): " + ", ".join(missing))
    return re.sub(r"{{(\w+)}}", lambda m: str(fields[m.group(1)]), text)


def render(sections, fields):
    """Sections -> [(style, text)] with placeholders filled."""
    out = []
    for heading, paras in sections:
        if heading not in UNPRINTED_HEADINGS:
            out.append((HEADING_STYLE, heading))
        out.extend((BODY_STYLE, fill(p, fields)) for p in paras)
    return out


# --- Fields -------------------------------------------------------------------

def long_date(iso):
    d = date.fromisoformat(iso)
    return f"{d:%B} {d.day}, {d.year}"


def build_fields(intake, deadlines=None):
    missing = [k for k in REQUIRED_FIELDS if not str(intake.get(k, "")).strip()]
    if missing:
        raise KeyError("Intake is missing: " + ", ".join(missing))
    if intake["matter_type"] not in MATTER_TYPES:
        raise ValueError(f"matter_type must be one of {MATTER_TYPES}")
    fields = {k: str(v).strip() for k, v in intake.items() if isinstance(v, (str, int, float))}
    fields.setdefault("client_address", "[Client Address]")
    fields["letter_date"] = long_date(intake.get("letter_date") or date.today().isoformat())
    if deadlines:
        fields.update({
            "service_date": long_date(deadlines["service_date"]),
            "service_method": SERVICE_METHOD_LABELS[deadlines["service_method"]],
            "answer_deadline": long_date(deadlines["answer_deadline"]),
            "signing_deadline": long_date(deadlines["signing_deadline"]),
        })
    return fields


def letter_paragraphs(intake, clauses_path, deadlines=None):
    """The tailored letter body, as [(style, text)]."""
    fields = build_fields(intake, deadlines)
    standard = parse_sections(STANDARD_FORM.read_text(encoding="utf-8"))
    clauses = parse_sections(Path(clauses_path).read_text(encoding="utf-8"))
    return render(merge_sections(standard, clauses), fields)


def standard_paragraphs(intake):
    """The firm's standard form with the same intake values, as [(style, text)]."""
    fields = build_fields(intake)
    return render(parse_sections(STANDARD_FORM.read_text(encoding="utf-8")), fields)


# --- Word document ------------------------------------------------------------

def set_text(paragraph, text):
    """Replace a paragraph's text, keeping the first run's formatting; '\\n' becomes a line break."""
    runs = paragraph.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    run = runs[0] if runs else paragraph.add_run()
    run.text = ""
    for i, line in enumerate(text.split("\n")):
        if i:
            run.add_break()
        run.add_text(line)


def open_letterhead(fields):
    """Open the template, fill the letterhead placeholders, and return (doc, body_placeholder)."""
    doc = Document(str(TEMPLATE))
    body = None
    for p in doc.paragraphs:
        if p.text.strip() == "{{BODY}}":
            body = p
        elif "{{" in p.text:
            set_text(p, fill(p.text, fields))
    if body is None:
        raise ValueError("Template has no {{BODY}} paragraph.")
    return doc, body


def write_letter(paragraphs, fields, out_path):
    doc, body = open_letterhead(fields)
    for style, text in paragraphs:
        set_text(body.insert_paragraph_before(style=style), text)
    body._p.getparent().remove(body._p)
    doc.save(str(out_path))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--intake", required=True, help="intake JSON file")
    ap.add_argument("--clauses", required=True, help="clause file for the matter type")
    ap.add_argument("--deadlines", help="JSON from compute_deadline.py (litigation only)")
    ap.add_argument("--out", required=True, help="output .docx path")
    args = ap.parse_args(argv)

    intake = json.loads(Path(args.intake).read_text(encoding="utf-8"))
    deadlines = json.loads(Path(args.deadlines).read_text(encoding="utf-8")) if args.deadlines else None
    matter = intake.get("matter_type")
    if matter == "litigation" and not deadlines:
        sys.exit("Litigation letters need --deadlines from compute_deadline.py.")
    if matter != "litigation" and deadlines:
        sys.exit("--deadlines is only used for litigation matters.")
    if Path(args.clauses).name != f"clauses-{matter}.md":
        sys.exit(f"matter_type is {matter!r} but the clause file is {Path(args.clauses).name!r}.")

    try:
        paragraphs = letter_paragraphs(intake, args.clauses, deadlines)
        write_letter(paragraphs, build_fields(intake, deadlines), args.out)
    except (KeyError, ValueError) as e:
        sys.exit(f"Error: {e.args[0]}")
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
